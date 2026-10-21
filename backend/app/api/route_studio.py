"""Validated PDF Maker API endpoints and lightweight processing history."""
import asyncio
import io
import html
import mimetypes
import json
import os
import re
import shutil
import sqlite3
import subprocess
import tempfile
import time
import uuid
import zipfile
from datetime import datetime, timezone
from pathlib import Path

import pymupdf
from fastapi import APIRouter, File, Form, HTTPException, Query, UploadFile
from fastapi.responses import FileResponse, JSONResponse, Response
from PIL import Image, UnidentifiedImageError
from pypdf import PdfReader, PdfWriter
from docx import Document

router = APIRouter(prefix="/api", tags=["PDF Maker"])
ROOT = Path(__file__).resolve().parents[2]
ON_VERCEL = bool(os.getenv("VERCEL"))
DEFAULT_STORE = Path(tempfile.gettempdir()) / "pdf_maker_temp" if ON_VERCEL else ROOT / "temp_storage"
STORE = Path(os.getenv("PDF_STORAGE_DIR", str(DEFAULT_STORE))).resolve()
STORE.mkdir(parents=True, exist_ok=True)
DB = Path(os.getenv("PDF_DATABASE_PATH", str(STORE / "history.sqlite3")))
MAX_BYTES = int(os.getenv("MAX_UPLOAD_SIZE", str(50 * 1024 * 1024)))
if ON_VERCEL:
    # Vercel Functions cap request bodies at 4.5 MB; leave room for multipart framing.
    MAX_BYTES = min(MAX_BYTES, 4 * 1024 * 1024)
RETENTION = int(os.getenv("FILE_RETENTION_HOURS", "24")) * 3600
MAX_PAGES = int(os.getenv("MAX_PDF_PAGES", "300"))
MAX_RENDER_DPI = int(os.getenv("MAX_RENDER_DPI", "300"))


def _db():
    DB.parent.mkdir(parents=True, exist_ok=True)
    with sqlite3.connect(DB) as conn:
        conn.execute("""CREATE TABLE IF NOT EXISTS jobs (
          id TEXT PRIMARY KEY, operation TEXT NOT NULL, created_at TEXT NOT NULL,
          status TEXT NOT NULL, input_size INTEGER NOT NULL, output_size INTEGER NOT NULL,
          file_id TEXT)""")


def _cleanup():
    cutoff = time.time() - RETENTION
    for p in STORE.iterdir():
        if p.is_file() and p.name != DB.name:
            try:
                if p.stat().st_mtime < cutoff:
                    p.unlink()
            except OSError:
                pass


async def _read(file: UploadFile, kind: str) -> bytes:
    name = (file.filename or "").lower()
    allowed = {"pdf": (".pdf",), "docx": (".docx",), "image": (".png", ".jpg", ".jpeg")}
    if not name.endswith(allowed[kind]):
        raise HTTPException(415, f"Unsupported file extension. Expected {', '.join(allowed[kind])}.")
    mime_allowed = {"pdf": {"application/pdf", "application/octet-stream"},
                    "docx": {"application/vnd.openxmlformats-officedocument.wordprocessingml.document", "application/octet-stream"},
                    "image": {"image/png", "image/jpeg", "application/octet-stream"}}
    if file.content_type and file.content_type not in mime_allowed[kind]:
        raise HTTPException(415, "The declared file type does not match a supported upload type.")
    data = await file.read(MAX_BYTES + 1)
    if len(data) > MAX_BYTES:
        raise HTTPException(413, f"File exceeds the {MAX_BYTES // (1024 * 1024)} MB upload limit.")
    if kind == "pdf" and not data.startswith(b"%PDF-"):
        raise HTTPException(400, "The uploaded file is not a valid PDF.")
    if kind == "docx" and not data.startswith(b"PK"):
        raise HTTPException(400, "The uploaded file is not a valid DOCX document.")
    if kind == "image":
        try:
            with Image.open(io.BytesIO(data)) as im:
                im.verify()
        except (UnidentifiedImageError, OSError):
            raise HTTPException(400, "The uploaded file is not a valid PNG or JPEG image.")
    return data


async def _read_many(files, kind: str):
    results = []
    total = 0
    for file in files:
        data = await _read(file, kind)
        total += len(data)
        if total > MAX_BYTES:
            raise HTTPException(413, f"Combined uploads exceed the {MAX_BYTES // (1024 * 1024)} MB request limit.")
        results.append(data)
    return results


def _pdf(data: bytes) -> PdfReader:
    try:
        reader = PdfReader(io.BytesIO(data), strict=True)
        if reader.is_encrypted:
            raise HTTPException(400, "Encrypted PDFs are not supported.")
        if not reader.pages:
            raise HTTPException(400, "PDF contains no pages.")
        if len(reader.pages) > MAX_PAGES:
            raise HTTPException(413, f"PDF exceeds the {MAX_PAGES}-page limit.")
        return reader
    except HTTPException:
        raise
    except Exception:
        raise HTTPException(400, "Unable to parse the uploaded PDF.")


def _result(operation: str, input_size: int, output: bytes, ext: str, filename: str, details=None):
    _cleanup()
    if ON_VERCEL and len(output) > 4 * 1024 * 1024:
        raise HTTPException(413, "The generated file exceeds Vercel's 4 MB safe response limit. Use the Docker deployment for larger documents.")
    file_id, job_id = (None if ON_VERCEL else uuid.uuid4().hex), uuid.uuid4().hex
    path = STORE / f"{file_id}{ext}" if file_id else None
    if path:
        path.write_bytes(output)
    _db()
    with sqlite3.connect(DB) as conn:
        conn.execute("INSERT INTO jobs VALUES (?, ?, ?, ?, ?, ?, ?)", (job_id, operation,
          datetime.now(timezone.utc).isoformat(), "completed", input_size, len(output), file_id))
    response = {"job_id": job_id, "file_id": file_id, "filename": filename,
                         "input_size": input_size, "output_size": len(output),
                         "reduction_percent": round((input_size - len(output)) * 100 / input_size, 2) if input_size else 0,
                         "download_url": f"/api/download/{file_id}"}
    if details: response.update(details)
    if ON_VERCEL:
        headers = {"Content-Disposition": f'attachment; filename="{filename}"', "X-Job-ID": job_id}
        if details:
            engine = details.get("conversion_engine")
            notice = details.get("conversion_notice")
            if engine: headers["X-Conversion-Engine"] = engine
            if notice: headers["X-Conversion-Notice"] = notice
        return Response(content=output, media_type=mimetypes.guess_type(filename)[0] or "application/octet-stream", headers=headers)
    return JSONResponse(response)


def _serialize(writer: PdfWriter) -> bytes:
    buf = io.BytesIO()
    writer.write(buf)
    return buf.getvalue()


def _merge_bytes(contents):
    writer = PdfWriter()
    for data in contents:
        for page in _pdf(data).pages:
            writer.add_page(page)
    return _serialize(writer)


def _split_bytes(data, wanted):
    reader = _pdf(data); writer = PdfWriter()
    for index in wanted: writer.add_page(reader.pages[index])
    return _serialize(writer)


def _compress_bytes(data, level):
    doc = pymupdf.open(stream=data, filetype="pdf")
    try: return doc.tobytes(garbage=level, deflate=True, deflate_images=True, deflate_fonts=True)
    finally: doc.close()


def _to_docx_bytes(data):
    doc = pymupdf.open(stream=data, filetype="pdf")
    word = Document()
    for i, page in enumerate(doc):
        if i: word.add_page_break()
        for paragraph in page.get_text("text").split("\n\n"):
            if paragraph.strip(): word.add_paragraph(paragraph.strip())
    doc.close(); out = io.BytesIO(); word.save(out)
    return out.getvalue()


def _extract_text_bytes(data):
    doc = pymupdf.open(stream=data, filetype="pdf")
    try: return "\n\n".join(f"--- Page {i + 1} ---\n{page.get_text('text')}" for i, page in enumerate(doc)).encode("utf-8")
    finally: doc.close()


def _render_pages(data, indices, dpi):
    doc = pymupdf.open(stream=data, filetype="pdf")
    mem = io.BytesIO()
    try:
        with zipfile.ZipFile(mem, "w", zipfile.ZIP_DEFLATED) as archive:
            for i in indices:
                rect = doc[i].rect
                if rect.width * dpi / 72 > 12000 or rect.height * dpi / 72 > 12000:
                    raise HTTPException(413, "A page is too large at the requested resolution. Lower DPI or use another page size.")
                pix = doc[i].get_pixmap(dpi=dpi, alpha=False)
                archive.writestr(f"page-{i+1:03}.png", pix.tobytes("png"))
        return mem.getvalue()
    finally: doc.close()


def _images_pdf(contents):
    writer = PdfWriter()
    for data in contents:
        with Image.open(io.BytesIO(data)) as im:
            if im.width * im.height > 40_000_000: raise HTTPException(413, "Image dimensions exceed the 40 megapixel limit.")
            rgb = im.convert("RGB"); buf = io.BytesIO(); rgb.save(buf, "PDF", resolution=150)
        writer.add_page(PdfReader(buf).pages[0])
    return _serialize(writer)


def _basic_docx_to_pdf(data: bytes) -> bytes:
    """Portable fallback for serverless hosts without LibreOffice; text and tables only."""
    from docx import Document as WordDocument
    from docx.oxml.text.paragraph import CT_P
    from docx.oxml.table import CT_Tbl
    from docx.text.paragraph import Paragraph
    from docx.table import Table as WordTable
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.enums import TA_LEFT
    from reportlab.platypus import SimpleDocTemplate, Paragraph as PdfParagraph, Spacer, Table, TableStyle

    word = WordDocument(io.BytesIO(data))
    styles = getSampleStyleSheet()
    normal = styles["BodyText"]
    normal.fontName = "Helvetica"
    normal.fontSize = 10
    normal.leading = 14
    heading = ParagraphStyle("DocHeading", parent=styles["Heading2"], fontName="Helvetica-Bold", alignment=TA_LEFT)
    story = []
    for block in word.element.body.iterchildren():
        if isinstance(block, CT_P):
            paragraph = Paragraph(block, word)
            pieces = []
            for run in paragraph.runs:
                value = html.escape(run.text).replace("\n", "<br/>")
                if run.bold: value = f"<b>{value}</b>"
                if run.italic: value = f"<i>{value}</i>"
                if run.underline: value = f"<u>{value}</u>"
                pieces.append(value)
            text = "".join(pieces).strip()
            if not text:
                continue
            is_heading = paragraph.style and paragraph.style.name.lower().startswith("heading")
            story.append(PdfParagraph(text, heading if is_heading else normal))
            story.append(Spacer(1, 5))
        elif isinstance(block, CT_Tbl):
            table = WordTable(block, word)
            rows = [[PdfParagraph(html.escape(cell.text).replace("\n", "<br/>"), normal) for cell in row.cells]
                    for row in table.rows]
            if rows:
                pdf_table = Table(rows, repeatRows=1, hAlign="LEFT")
                pdf_table.setStyle(TableStyle([
                    ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#9ca3af")),
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#eef2f7")),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("LEFTPADDING", (0, 0), (-1, -1), 6),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                ]))
                story.extend((pdf_table, Spacer(1, 8)))
    out = io.BytesIO()
    SimpleDocTemplate(out, pagesize=A4, rightMargin=48, leftMargin=48, topMargin=48, bottomMargin=48,
                      title=word.core_properties.title or "Converted document",
                      author=word.core_properties.author or "").build(story)
    return out.getvalue()


def _find_libreoffice():
    candidates = [os.getenv("LIBREOFFICE_PATH"), shutil.which("soffice"), shutil.which("libreoffice")]
    if os.name == "nt":
        candidates.extend((r"C:\Program Files\LibreOffice\program\soffice.exe",
                           r"C:\Program Files (x86)\LibreOffice\program\soffice.exe"))
    for candidate in candidates:
        if candidate and (Path(candidate).is_file() or shutil.which(candidate)):
            return str(candidate)
    return None


@router.get("/health")
async def health():
    return {"status": "ok", "service": "PDF Maker"}


@router.get("/tools")
async def tools():
    return {"tools": ["merge", "split", "compress", "pdf-to-docx", "docx-to-pdf", "extract-text", "pdf-to-images", "images-to-pdf", "metadata"]}


@router.post("/pdf/merge")
async def merge(files: list[UploadFile] = File(...)):
    if not 2 <= len(files) <= 20:
        raise HTTPException(400, "Upload between 2 and 20 PDFs, in the order to merge.")
    contents = await _read_many(files, "pdf")
    try:
        output = await asyncio.to_thread(_merge_bytes, contents)
    except HTTPException:
        raise
    return _result("merge", sum(map(len, contents)), output, ".pdf", "merged.pdf")


@router.post("/pdf/split")
async def split(file: UploadFile = File(...), pages: str = Form(..., description="Page ranges, e.g. 1-3,5")):
    data = await _read(file, "pdf"); reader = _pdf(data)
    wanted = []
    try:
        for part in pages.split(","):
            match = re.fullmatch(r"\s*(\d+)\s*(?:-\s*(\d+)\s*)?", part)
            if not match: raise ValueError
            start, end = int(match.group(1)), int(match.group(2) or match.group(1))
            if start < 1 or end < start or end > len(reader.pages): raise ValueError
            wanted.extend(range(start - 1, end))
        if not wanted or len(set(wanted)) != len(wanted): raise ValueError
    except ValueError:
        raise HTTPException(422, f"Invalid page ranges. Use 1-based ranges within 1-{len(reader.pages)}, such as 1-3,5.")
    output = await asyncio.to_thread(_split_bytes, data, wanted)
    return _result("split", len(data), output, ".pdf", "split.pdf")


@router.post("/pdf/compress")
async def compress(file: UploadFile = File(...), level: int = Form(2, ge=1, le=3)):
    data = await _read(file, "pdf"); _pdf(data)
    try:
        output = await asyncio.to_thread(_compress_bytes, data, level)
    except Exception:
        raise HTTPException(400, "Unable to optimize this PDF.")
    result = _result("compress", len(data), output, ".pdf", "compressed.pdf")
    # Compression accounting is returned in job metadata and the response.
    return result


@router.post("/pdf/to-docx")
async def to_docx(file: UploadFile = File(...)):
    data = await _read(file, "pdf"); _pdf(data)
    output = await asyncio.to_thread(_to_docx_bytes, data)
    return _result("pdf-to-docx", len(data), output, ".docx", "converted.docx")


@router.post("/docx/to-pdf")
async def docx_to_pdf(file: UploadFile = File(...)):
    data = await _read(file, "docx")
    office = _find_libreoffice()
    if office:
        with tempfile.TemporaryDirectory(dir=STORE) as tmp:
            src = Path(tmp) / "input.docx"; src.write_bytes(data)
            try:
                proc = await asyncio.to_thread(subprocess.run, [office, "--headless", "--convert-to", "pdf", "--outdir", tmp, str(src)],
                    capture_output=True, timeout=90, check=False)
            except (OSError, subprocess.TimeoutExpired):
                raise HTTPException(503, "LibreOffice conversion failed or timed out.")
            outpath = Path(tmp) / "input.pdf"
            if proc.returncode or not outpath.is_file():
                raise HTTPException(422, "LibreOffice could not convert this DOCX file.")
            output = outpath.read_bytes()
        engine = "LibreOffice"
        notice = None
    else:
        if os.getenv("ALLOW_BASIC_DOCX_FALLBACK", "false").strip().lower() != "true":
            raise HTTPException(503, "Format-preserving DOCX-to-PDF conversion requires LibreOffice. Vercel does not provide it; deploy the API with the included Dockerfile or connect a LibreOffice-enabled backend. The altered-layout fallback is disabled by default.")
        try:
            output = await asyncio.to_thread(_basic_docx_to_pdf, data)
        except Exception:
            raise HTTPException(422, "The DOCX could not be converted by the basic renderer. Try LibreOffice for broader format support.")
        engine = "ReportLab basic renderer"
        notice = "Basic conversion: text and tables only. Complex layout, images, columns, and exact formatting may differ."
    return _result("docx-to-pdf", len(data), output, ".pdf", "converted.pdf",
                   {"conversion_engine": engine, "conversion_notice": notice} if notice else {"conversion_engine": engine})


@router.post("/pdf/extract-text")
async def extract_text(file: UploadFile = File(...)):
    data = await _read(file, "pdf"); _pdf(data)
    text = await asyncio.to_thread(_extract_text_bytes, data)
    return _result("extract-text", len(data), text, ".txt", "extracted.txt")


@router.post("/pdf/to-images")
async def to_images(file: UploadFile = File(...), pages: str = Form("all"), dpi: int = Form(150, ge=72, le=300)):
    if dpi > MAX_RENDER_DPI: raise HTTPException(422, f"Maximum render resolution is {MAX_RENDER_DPI} DPI.")
    data = await _read(file, "pdf"); _pdf(data)
    doc = pymupdf.open(stream=data, filetype="pdf")
    try:
        indices = list(range(len(doc))) if pages == "all" else []
        if pages != "all":
            for part in pages.split(","):
                m = re.fullmatch(r"\s*(\d+)(?:-(\d+))?\s*", part)
                if not m: raise ValueError
                a, b = int(m.group(1)), int(m.group(2) or m.group(1))
                if a < 1 or b < a or b > len(doc): raise ValueError
                indices.extend(range(a-1, b))
        if not indices or len(indices) > 50: raise ValueError
    except ValueError:
        doc.close(); raise HTTPException(422, "Invalid page range; choose up to 50 valid pages.")
    doc.close()
    output = await asyncio.to_thread(_render_pages, data, indices, dpi)
    return _result("pdf-to-images", len(data), output, ".zip", "pages.zip")


@router.post("/images/to-pdf")
async def images_to_pdf(files: list[UploadFile] = File(...)):
    if not 1 <= len(files) <= 20: raise HTTPException(400, "Upload between 1 and 20 images.")
    contents = await _read_many(files, "image")
    total = sum(map(len, contents))
    try: output = await asyncio.to_thread(_images_pdf, contents)
    except HTTPException: raise
    except Exception: raise HTTPException(400, "Unable to process one of the uploaded images.")
    return _result("images-to-pdf", total, output, ".pdf", "images.pdf")


@router.get("/pdf/metadata")
async def metadata_lookup(file_id: str = Query(..., pattern=r"^[a-f0-9]{32}$")):
    path = STORE / f"{file_id}.pdf"
    if not path.is_file() or path.resolve().parent != STORE: raise HTTPException(404, "PDF not found or expired.")
    return _metadata(path.read_bytes())


def _metadata(data: bytes):
    reader = _pdf(data); info = reader.metadata
    return {"page_count": len(reader.pages), "title": getattr(info, "title", None),
            "author": getattr(info, "author", None), "creation_date": str(getattr(info, "creation_date", "") or "") or None,
            "file_size": len(data)}


@router.post("/pdf/metadata")
async def metadata_upload(file: UploadFile = File(...)):
    data = await _read(file, "pdf")
    return _metadata(data)


@router.get("/jobs")
async def jobs(limit: int = Query(20, ge=1, le=100)):
    _db()
    with sqlite3.connect(DB) as conn:
        conn.row_factory = sqlite3.Row
        return {"jobs": [dict(row) for row in conn.execute("SELECT id AS job_id, operation, created_at, status, input_size, output_size FROM jobs ORDER BY created_at DESC LIMIT ?", (limit,))]}


@router.get("/jobs/{job_id}")
async def job(job_id: str):
    _db()
    with sqlite3.connect(DB) as conn:
        conn.row_factory = sqlite3.Row
        row = conn.execute("SELECT id AS job_id, operation, created_at, status, input_size, output_size FROM jobs WHERE id=?", (job_id,)).fetchone()
    if not row: raise HTTPException(404, "Job not found.")
    return dict(row)


@router.get("/download/{file_id}")
async def download(file_id: str):
    if ON_VERCEL:
        raise HTTPException(404, "Vercel downloads are returned directly by the processing request; the temporary filesystem is not shared between requests.")
    if not re.fullmatch(r"[a-f0-9]{32}", file_id): raise HTTPException(404, "Download not found or expired.")
    matches = list(STORE.glob(f"{file_id}.*"))
    if not matches or matches[0].resolve().parent != STORE: raise HTTPException(404, "Download not found or expired.")
    return FileResponse(matches[0], filename=f"download{matches[0].suffix}")
